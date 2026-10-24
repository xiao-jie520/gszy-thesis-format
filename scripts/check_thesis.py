# -*- coding: utf-8 -*-
"""
check_thesis.py — 按《甘肃中医药大学 医学信息工程学院本科毕业论文（设计）撰写规范》
体检一份 .docx，输出 Markdown 检查报告。

用法:
    python check_thesis.py 论文.docx
    python check_thesis.py 论文.docx -o 检查报告.md

退出码: 0 无错误（可能有警告）；1 存在错误；2 无法解析文件。
依赖: python-docx
"""
import argparse
import re
import sys
from collections import OrderedDict

try:
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph
except ImportError:  # pragma: no cover
    sys.exit("缺少依赖 python-docx，请先执行: pip install python-docx")

TOL = 0.08           # cm / pt 容差
A4 = (21.0, 29.7)
MARGINS = {"上": 3.5, "下": 2.5, "左": 2.5, "右": 2.5}
HEADER_DIST, FOOTER_DIST = 2.5, 2.0
HEADER_KEY = "甘肃中医药大学本科毕业论文（设计）"

CAP_FIG = re.compile(r"^图\s*[A-Z]?\d+\s*[-－.]\s*\d+")
CAP_TAB = re.compile(r"^表\s*[A-Z]?\d+\s*[-－.]\s*\d+")
REF = re.compile(r"^\[\d+\]")
CHAPTER = re.compile(r"^第\s*(\d+)\s*章")

errors, warnings, infos = [], [], []


def err(msg):
    errors.append(msg)


def warn(msg):
    warnings.append(msg)


def info(msg):
    infos.append(msg)


# ---------------------------------------------------------------- 属性解析
def _rfonts(rpr):
    if rpr is None:
        return None
    return rpr.find(qn("w:rFonts"))


def eff_eastasia(run, para):
    rf = _rfonts(run._element.rPr)
    if rf is not None and rf.get(qn("w:eastAsia")):
        return rf.get(qn("w:eastAsia"))
    st = para.style
    while st is not None:
        rf = _rfonts(st.element.rPr)
        if rf is not None and rf.get(qn("w:eastAsia")):
            return rf.get(qn("w:eastAsia"))
        st = st.base_style
    return "宋体"


def eff_ascii(run, para):
    rf = _rfonts(run._element.rPr)
    if rf is not None and rf.get(qn("w:ascii")):
        return rf.get(qn("w:ascii"))
    st = para.style
    while st is not None:
        rf = _rfonts(st.element.rPr)
        if rf is not None and rf.get(qn("w:ascii")):
            return rf.get(qn("w:ascii"))
        st = st.base_style
    return "Times New Roman"


def eff_size(run, para):
    if run.font.size is not None:
        return round(run.font.size.pt, 1)
    st = para.style
    while st is not None:
        if st.font.size is not None:
            return round(st.font.size.pt, 1)
        st = st.base_style
    return 12.0


def first_line_chars(para):
    ind = para._p.pPr.find(qn("w:ind")) if para._p.pPr is not None else None
    if ind is not None:
        v = ind.get(qn("w:firstLineChars"))
        if v:
            return int(v) / 100.0
    fli = para.paragraph_format.first_line_indent
    if fli is None:
        return 0.0
    return round(fli.pt / 12.0, 2)


def line_spacing(para):
    ls = para.paragraph_format.line_spacing
    if ls is None:
        st = para.style
        while st is not None:
            if st.paragraph_format.line_spacing is not None:
                ls = st.paragraph_format.line_spacing
                break
            st = st.base_style
    return ls


def head_text(p):
    return re.sub(r"[\s\u3000]", "", p.text)


class Checker:
    def __init__(self, path):
        self.doc = docx.Document(path)
        self.path = path

    # ---------------------------------------------------------- 1 页面
    def check_page(self):
        for i, s in enumerate(self.doc.sections, 1):
            if abs(s.page_width.cm - A4[0]) > TOL or abs(s.page_height.cm - A4[1]) > TOL:
                err(f"第{i}节纸张不是 A4：{s.page_width.cm:.1f}×{s.page_height.cm:.1f} cm")
            for name, got, exp in (
                ("上", s.top_margin.cm, MARGINS["上"]),
                ("下", s.bottom_margin.cm, MARGINS["下"]),
                ("左", s.left_margin.cm, MARGINS["左"]),
                ("右", s.right_margin.cm, MARGINS["右"]),
            ):
                if abs(got - exp) > TOL:
                    err(f"第{i}节{name}边距 {got:.2f} cm，应为 {exp} cm")
            if abs(s.header_distance.cm - HEADER_DIST) > TOL:
                warn(f"第{i}节页眉距边界 {s.header_distance.cm:.2f} cm，规范为 {HEADER_DIST} cm")
            if abs(s.footer_distance.cm - FOOTER_DIST) > TOL:
                warn(f"第{i}节页脚距边界 {s.footer_distance.cm:.2f} cm，规范为 {FOOTER_DIST} cm")

    # ---------------------------------------------------------- 2 页码 / 页眉
    def check_numbering(self):
        fmts = []
        for s in self.doc.sections:
            pn = s._sectPr.find(qn("w:pgNumType"))
            fmt = pn.get(qn("w:fmt")) if pn is not None else None
            fmts.append(fmt)
        if "upperRoman" not in fmts:
            warn("未发现前置部分（摘要/Abstract/目录）使用罗马数字页码")
        if fmts and fmts[-1] not in ("decimal", None):
            warn(f"最后一节的页码格式是 {fmts[-1]}，正文应为阿拉伯数字 decimal")

        has_page_field = False
        for s in self.doc.sections:
            xml = (s.footer._element.xml if s.footer is not None else "")
            if "PAGE" in xml:
                has_page_field = True
                break
        if not has_page_field:
            warn("页脚中未检测到页码域（PAGE）")

        header_hit = 0
        for s in self.doc.sections:
            t = re.sub(r"[\s\u3000]", "", s.header.paragraphs[0].text) if s.header.paragraphs else ""
            if HEADER_KEY.replace("（", "").replace("）", "") in t.replace("（", "").replace("）", ""):
                header_hit += 1
        if header_hit == 0:
            warn(f"页眉未找到“{HEADER_KEY}”")
        elif header_hit == 1:
            info("页眉仅出现在 1 个节（通常正文节有页眉、前置部分留空即可）")

    # ---------------------------------------------------------- 3 标题
    def check_headings(self):
        h1 = h2 = h3 = h4 = 0
        for p in self.doc.paragraphs:
            name = p.style.name
            if name == "Heading 4" or name == "Heading 5":
                h4 += 1
                err(f"出现四级标题「{p.text.strip()[:20]}」，规范最多三级")
                continue
            sizes = {eff_size(r, p) for r in p.runs if r.text.strip()}
            fonts = {eff_eastasia(r, p) for r in p.runs if r.text.strip()}
            if name == "Heading 1":
                h1 += 1
                if sizes and 15.0 not in sizes:
                    err(f"一级标题字号 {sorted(sizes)}，应为 15 pt（小三）：{p.text.strip()[:20]}")
                if fonts and "黑体" not in fonts:
                    err(f"一级标题字体 {sorted(fonts)}，应为黑体：{p.text.strip()[:20]}")
                if p.alignment != WD_ALIGN_PARAGRAPH.CENTER:
                    warn(f"一级标题未居中：{p.text.strip()[:20]}")
            elif name in ("Heading 2", "Heading 3"):
                h2 = h2 + 1 if name == "Heading 2" else h2
                h3 = h3 + 1 if name == "Heading 3" else h3
                if sizes and 12.0 not in sizes:
                    err(f"{name} 字号 {sorted(sizes)}，应为 12 pt（小四）：{p.text.strip()[:20]}")
                if fonts and "黑体" not in fonts:
                    err(f"{name} 字体 {sorted(fonts)}，应为黑体：{p.text.strip()[:20]}")
        if h1 == 0:
            err("未检测到任何一级标题（章标题）")
        info(f"标题统计：一级 {h1} 个，二级 {h2} 个，三级 {h3} 个，四级 {h4} 个")

    # ---------------------------------------------------------- 4 正文
    def check_body(self):
        bad_font, bad_size, bad_space, bad_indent, checked = [], [], [], [], 0
        started = False
        for i, p in enumerate(self.doc.paragraphs):
            txt = p.text.strip()
            if not started:
                if head_text(p) in ("摘要", "Abstract"):
                    started = True
                continue
            if not txt:
                continue
            style = p.style.name.lower()
            if style.startswith("heading") or style.startswith("toc"):
                continue
            if p._p.findall(".//" + qn("w:instrText")):
                continue          # 域代码（目录占位）
            if REF.match(txt) or CAP_FIG.match(txt) or CAP_TAB.match(txt):
                continue
            if txt.startswith("【关键词】") or txt.lower().startswith("【keywords】"):
                continue
            if p._p.findall(".//" + qn("w:drawing")) or p._p.findall(".//" + qn("w:pict")):
                continue
            if any(eff_ascii(r, p) in ("Consolas", "Courier New") for r in p.runs):
                continue          # 附录代码
            checked += 1
            sizes = {eff_size(r, p) for r in p.runs if r.text.strip()}
            eas = {eff_eastasia(r, p) for r in p.runs if r.text.strip()}
            if sizes and 12.0 not in sizes:
                bad_size.append((i, txt[:18], sorted(sizes)))
            if eas and "宋体" not in eas:
                bad_font.append((i, txt[:18], sorted(eas)))
            ls = line_spacing(p)
            if ls is not None and abs(float(ls) - 1.5) > 0.01:
                bad_space.append((i, txt[:18], ls))
            if abs(first_line_chars(p) - 2.0) > 0.05:
                bad_indent.append((i, txt[:18], first_line_chars(p)))

        for label, lst, expect in (
            ("字号不是小四(12pt)", bad_size, "12 pt"),
            ("中文字体不是宋体", bad_font, "宋体"),
            ("行距不是 1.5 倍", bad_space, "1.5"),
            ("首行缩进不是 2 字符", bad_indent, "2 字符"),
        ):
            if lst:
                sample = "；".join(f"第{a}段「{b}」→{c}" for a, b, c in lst[:4])
                msg = f"{label}：共 {len(lst)}/{checked} 段（如 {sample}）"
                (err if label in ("字号不是小四(12pt)", "中文字体不是宋体") else warn)(msg)
        info(f"正文段落检查了 {checked} 段")

    # ---------------------------------------------------------- 5 图表
    def check_floats(self):
        body = list(self.doc.element.body.iterchildren())
        blocks = []
        for el in body:
            if el.tag == qn("w:p"):
                blocks.append(Paragraph(el, self.doc))
            elif el.tag == qn("w:tbl"):
                blocks.append(Table(el, self.doc))
            else:
                blocks.append(el)
        fig_nums, tab_nums = [], []
        for idx, blk in enumerate(blocks):
            if not isinstance(blk, Paragraph):
                continue
            t = blk.text.strip()
            if not t:
                continue
            if CAP_FIG.match(t):
                fig_nums.append(t.split()[0])
                prev = blocks[idx - 1] if idx else None
                if isinstance(prev, Paragraph) and prev.text.strip():
                    warn(f"图题上方应紧接图片，此处上方是文字：{t[:20]}")
                if not re.match(r"^图\s*\d+\s*[-－.]\s*\d+", t):
                    warn(f"图号未按“图章-序”编号：{t[:20]}")
            if CAP_TAB.match(t):
                tab_nums.append(t.split()[0])
                if not re.match(r"^表\s*\d+\s*[-－.]\s*\d+", t):
                    warn(f"表号未按“表章-序”编号：{t[:20]}")
                nxt = blocks[idx + 1] if idx + 1 < len(blocks) else None
                if not isinstance(nxt, Table):
                    warn(f"表题下方未紧跟表格：{t[:20]}")
        for label, nums, pat in (("图", fig_nums, "图"), ("表", tab_nums, "表")):
            chs = [int(m.group(1)) for n in nums if (m := re.match(pat + r"\s*(\d+)", n))]
            if chs != sorted(chs):
                warn(f"{label}编号未按章递增：{nums}")

    # ---------------------------------------------------------- 6 参考文献
    def check_refs(self):
        in_refs = False
        bad = 0
        total = 0
        for p in self.doc.paragraphs:
            if p.style.name == "Heading 1" and head_text(p) == "参考文献":
                in_refs = True
                continue
            if in_refs and p.style.name == "Heading 1":
                break
            t = p.text.strip()
            if in_refs and REF.match(t):
                total += 1
                sizes = {eff_size(r, p) for r in p.runs if r.text.strip()}
                if sizes and sizes != {10.5}:
                    bad += 1
        if in_refs is False:
            err("未找到“参考文献”章节")
        elif total == 0:
            err("参考文献章节下没有 [1] [2] … 条目")
        elif bad:
            warn(f"参考文献有 {bad}/{total} 条字号不是五号(10.5pt)")
        else:
            info(f"参考文献 {total} 条，字号正常")
        if in_refs:
            for p in self.doc.paragraphs:
                t = p.text.strip()
                if REF.match(t) and re.search(r"[\u4e00-\u9fff]，", t):
                    warn(f"参考文献条目中使用了中文逗号：{t[:24]}")
                    break

    # ---------------------------------------------------------- 7 关键词
    def check_keywords(self):
        for p in self.doc.paragraphs:
            t = p.text.strip()
            if t.startswith("【关键词】"):
                body = t[len("【关键词】"):].strip()
                items = [x for x in re.split(r"[;；]", body) if x.strip()]
                if not 3 <= len(items) <= 5:
                    err(f"中文关键词 {len(items)} 个，规范为 3～5 个")
                if body.rstrip().endswith(("；", ";", "，", "。")):
                    warn("中文关键词末尾不应加标点")
                if "，" in body:
                    warn("中文关键词之间应用分号“；”分隔，检测到逗号")
            if t.lower().startswith("【keywords】"):
                body = t[len("【Keywords】"):].strip()
                if "  " not in body:
                    warn("英文关键词之间应空两个空格")
                if body.rstrip().endswith((",", ".", ";")):
                    warn("英文关键词末尾不应加标点")

    # ---------------------------------------------------------- 8 标点与字体混用
    def check_punct(self):
        half_in_cn = []
        latin_wrong = []
        for i, p in enumerate(self.doc.paragraphs):
            t = p.text
            if not t.strip() or p.style.name.startswith("Heading"):
                continue
            if REF.match(t.strip()):
                continue          # 参考文献按规范用英文标点
            if re.search(r"[\u4e00-\u9fff][,;:!?]", t) or re.search(r"[,;:!?][\u4e00-\u9fff]", t):
                half_in_cn.append((i, t.strip()[:18]))
            for r in p.runs:
                if not r.text.strip():
                    continue
                if re.search(r"[A-Za-z0-9]", r.text) and eff_ascii(r, p) not in (
                        "Times New Roman", "Times New Roman PSMT", "Consolas", "Courier New"):
                    latin_wrong.append((i, r.text[:18], eff_ascii(r, p)))
                    break
        if half_in_cn:
            warn(f"中文语境中混用半角标点（{len(half_in_cn)} 段），如第{half_in_cn[0][0]}段「{half_in_cn[0][1]}」")
        if latin_wrong:
            warn(f"英文/数字所在文本字体不是 Times New Roman（{len(latin_wrong)} 段），如第{latin_wrong[0][0]}段「{latin_wrong[0][1]}」→{latin_wrong[0][2]}")

    # ---------------------------------------------------------- 9 空段与结构
    def check_structure(self):
        started = False
        blanks = 0
        for p in self.doc.paragraphs:
            if p.style.name == "Heading 1" and CHAPTER.match(p.text.strip()):
                started = True
            if not started or p.text.strip():
                continue
            if p._p.findall(".//" + qn("w:drawing")):
                continue
            if p._p.pPr is not None and p._p.pPr.find(qn("w:sectPr")) is not None:
                continue          # 分节符段落，不可见
            blanks += 1
        if blanks > 8:
            warn(f"正文有 {blanks} 个空段落，请用段落间距代替空行")
        order, missing = [], []
        for key, label in (("摘要", "中文摘要"), ("abstract", "Abstract"),
                           ("目录", "目录"), ("参考文献", "参考文献"),
                           ("致谢", "致谢")):
            found = any(head_text(p).lower() == key for p in self.doc.paragraphs)
            if not found:
                missing.append(label)
            else:
                order.append(label)
        if missing:
            warn("缺少必备部分：" + "、".join(missing))
        chapters = sum(1 for p in self.doc.paragraphs
                       if p.style.name == "Heading 1" and CHAPTER.match(p.text.strip()))
        if chapters == 0:
            warn("未发现“第N章”式章标题（若用“1 引言”式编号可忽略）")

    # ---------------------------------------------------------- 报告
    def report(self):
        self.check_page()
        self.check_numbering()
        self.check_headings()
        self.check_body()
        self.check_floats()
        self.check_refs()
        self.check_keywords()
        self.check_punct()
        self.check_structure()

        lines = ["# 论文格式检查报告", "", f"文件：`{self.path}`", ""]
        lines.append(f"**结果：{len(errors)} 项错误，{len(warnings)} 项警告**")
        lines.append("")
        if errors:
            lines += ["## 必须修改", ""]
            lines += [f"- ❌ {m}" for m in errors]
            lines.append("")
        if warnings:
            lines += ["## 建议核对", ""]
            lines += [f"- ⚠️ {m}" for m in warnings]
            lines.append("")
        if infos:
            lines += ["## 检查明细", ""]
            lines += [f"- {m}" for m in infos]
            lines.append("")
        if not errors and not warnings:
            lines += ["## 未发现明显格式问题 ✅", ""]
        return "\n".join(lines)


def main():
    ap = argparse.ArgumentParser(description="按学院论文规范检查 .docx 格式")
    ap.add_argument("docx", help="待检查的 .docx")
    ap.add_argument("-o", "--output", help="输出 Markdown 报告路径")
    args = ap.parse_args()

    try:
        c = Checker(args.docx)
        md = c.report()
    except Exception as e:  # noqa: BLE001
        print(f"无法解析文件：{e}", file=sys.stderr)
        return 2

    if args.output:
        with open(args.output, "w", encoding="utf-8") as f:
            f.write(md)
        print(f"报告已写入：{args.output}")
    print(md)
    return 1 if errors else 0


if __name__ == "__main__":
    sys.exit(main())
